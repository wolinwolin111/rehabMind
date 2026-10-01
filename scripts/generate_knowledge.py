"""Compile the project knowledge workbook into a validated runtime snapshot."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "knowledge" / "RehabMind_KnowledgeDB_v3.8.xlsx"
CLINICAL_EXTENSION = ROOT / "knowledge" / "RehabMind_ClinicalExtension_v1.1.xlsx"
OUTPUT = ROOT / "build" / "knowledge" / "runtime.json"
SHEETS = (
    "01_Region",
    "02_ClinicalItem",
    "03_Finding",
    "04_Intervention",
    "05_RegionClinicalMap",
    "06_FindingInterventionMap",
    "07_ContextRule",
    "08_AssessmentFindingMap",
    "09_Evidence",
    "10_Governance",
    "11_ContextItemMap",
)


def read_sheet(sheet: Any) -> list[dict[str, Any]]:
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)
    if not headers or not headers[0]:
        raise ValueError(f"Missing headers in {sheet.title}")
    result = []
    seen = set()
    for row in rows:
        key = row[0]
        if key is None or not str(key).strip():
            continue
        if key in seen:
            raise ValueError(f"Duplicate key in {sheet.title}: {key}")
        seen.add(key)
        item = {str(header): row[index] if index < len(row) else None
                for index, header in enumerate(headers) if header}
        result.append(item)
    return result


def validate(tables: dict[str, list[dict[str, Any]]]) -> None:
    def ids(sheet: str, field: str) -> set[str]:
        return {str(row[field]) for row in tables[sheet]}

    regions = ids("01_Region", "region_id")
    items = ids("02_ClinicalItem", "item_id")
    findings = ids("03_Finding", "finding_id")
    interventions = ids("04_Intervention", "intervention_id")
    rcm = tables["05_RegionClinicalMap"]
    afm = tables["08_AssessmentFindingMap"]
    fim = tables["06_FindingInterventionMap"]
    errors = []
    governance = {row["key"]: row["value"] for row in tables["10_Governance"]}
    modules = {part.split(":", 1)[0] for part in str(governance.get("module_labels", "")).split("|") if ":" in part}
    region_modules = {row["region_id"]: row["module"] for row in tables["01_Region"]}
    for row in tables["01_Region"]:
        if row["module"] not in modules:
            errors.append(f"Unregistered region module: {row['region_id']}")
    dimensions = set(str(governance.get("assessment_dimension_enum", "")).split("|"))
    modes = set(str(governance.get("display_mode_enum", "")).split("|"))
    for row in rcm:
        if row["region_id"] not in regions or row["clinical_item_id"] not in items:
            errors.append(f"Invalid region/item mapping: {row['map_id']}")
        if row["priority"] not in {"High", "Secondary", "Context"}:
            errors.append(f"Invalid priority: {row['map_id']}")
        if row.get("assessment_dimension") not in dimensions:
            errors.append(f"Invalid assessment dimension: {row['map_id']}")
        if not str(row.get("assessment_group") or "").strip():
            errors.append(f"Missing assessment group: {row['map_id']}")
        if not isinstance(row.get("display_order"), (int, float)):
            errors.append(f"Invalid display order: {row['map_id']}")
        if row.get("display_mode") not in modes:
            errors.append(f"Invalid display mode: {row['map_id']}")
    for row in afm:
        if row["clinical_item_id"] not in items or row["finding_id"] not in findings:
            errors.append(f"Invalid assessment/finding mapping: {row['map_id']}")
    for row in fim:
        if row["finding_id"] not in findings or row["intervention_id"] not in interventions:
            errors.append(f"Invalid finding/intervention mapping: {row['map_id']}")
    base_regions = {r["region_id"] for r in rcm if r["priority"] in {"High", "Secondary"}}
    mapped_items = {r["clinical_item_id"] for r in rcm}
    linked_items = {r["clinical_item_id"] for r in afm}
    reachable_findings = {r["finding_id"] for r in afm if r["clinical_item_id"] in mapped_items}
    linked_findings = {r["finding_id"] for r in fim}
    errors += [f"Region without base assessment: {rid}" for rid in sorted(regions - base_regions)]
    errors += [f"Mapped item without finding: {iid}" for iid in sorted(mapped_items - linked_items)]
    # Observation and screening findings may legitimately have no treatment.

    for rule in tables["07_ContextRule"]:
        if rule["rule_type"] == "PROFILE" and rule["target_type"] != "CLINICAL_CONSIDERATION":
            errors.append(f"Profile rule has invalid target: {rule['rule_id']}")
        scope = set(str(rule.get("module_codes") or "").split("|")) - {""}
        if not scope or not scope <= modules:
            errors.append(f"Invalid context module scope: {rule['rule_id']}")

    rules = {row["rule_id"]: row for row in tables["07_ContextRule"]}
    seen_links = set()
    for link in tables["11_ContextItemMap"]:
        rule = rules.get(link["rule_id"])
        if not rule or link["clinical_item_id"] not in items:
            errors.append(f"Invalid context item reference: {link['map_id']}")
            continue
        pair = (link["rule_id"], link["clinical_item_id"])
        if pair in seen_links:
            errors.append(f"Duplicate context item reference: {link['map_id']}")
        seen_links.add(pair)
        if rule["rule_type"] in {"PROFILE", "SAFETY"}:
            errors.append(f"Context item link has invalid rule type: {link['map_id']}")
        for module in str(rule["module_codes"]).split("|"):
            if not any(row["clinical_item_id"] == link["clinical_item_id"] and region_modules.get(row["region_id"]) == module for row in rcm):
                errors.append(f"Context target has no placement in {module}: {link['map_id']}")

    if not governance.get("database_version"):
        errors.append("Missing database version")
    if errors:
        raise ValueError("Knowledge validation failed:\n" + "\n".join(errors))


def main() -> None:
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    if any(name not in workbook.sheetnames for name in SHEETS):
        raise ValueError("Workbook is missing a required runtime or governance sheet")
    tables = {name: read_sheet(workbook[name]) for name in SHEETS}
    validate(tables)
    extension = load_workbook(CLINICAL_EXTENSION, read_only=True, data_only=True)
    required = ("DimensionGuide", "ClinicalPurpose", "FindingFeature", "InterventionReference", "ConsultationGuide", "LocationGuide")
    if any(name not in extension.sheetnames for name in required):
        raise ValueError("Clinical extension is missing a required sheet")
    extra = {name: read_sheet(extension[name]) for name in required}
    item_ids = {row["item_id"] for row in tables["02_ClinicalItem"]}
    dimensions = set(next(row["value"] for row in tables["10_Governance"] if row["key"] == "assessment_dimension_enum").split("|"))
    assessment_by_id = {row["map_id"]: row for row in tables["08_AssessmentFindingMap"]}
    assessment_ids = set(assessment_by_id)
    feature_ids = {row["Feature_ID"] for row in extra["FindingFeature"]}
    errors = []
    region_ids = {row["region_id"] for row in tables["01_Region"]}
    module_ids = {row["module"] for row in tables["01_Region"]}
    area_ids = {row["Area_ID"] for row in extra["LocationGuide"]}
    for row in extra["LocationGuide"]:
        if row.get('Auto_Apply_To_Region') is not None and not isinstance(row['Auto_Apply_To_Region'], bool):
            errors.append(f"Invalid location auto-apply flag: {row['Area_ID']}")
        ids = str(row.get("Region_IDs") or "").split("|")
        modules = set(str(row.get("Module_Codes") or "").split("|"))
        if not 1 <= len(ids) <= 3 or len(set(ids)) != len(ids) or not set(ids) <= region_ids:
            errors.append(f"Invalid location regions: {row['Area_ID']}")
        actual_modules = {r["module"] for r in tables["01_Region"] if r["region_id"] in ids}
        if modules != actual_modules:
            errors.append(f"Invalid location modules: {row['Area_ID']}")
        assigned = []
        for field in ("Primary_Item_IDs", "Additional_Item_IDs", "Linked_Item_IDs"):
            assigned += [item for item in str(row.get(field) or "").split("|") if item]
        if len(assigned) != len(set(assigned)) or not set(assigned) <= item_ids:
            errors.append(f"Invalid or duplicate location items: {row['Area_ID']}")
        for item in assigned:
            if not any(mapping['clinical_item_id'] == item and
                       any(r['region_id'] == mapping['region_id'] and r['module'] in modules for r in tables['01_Region'])
                       for mapping in tables['05_RegionClinicalMap']):
                errors.append(f"Location item lacks module placement: {row['Area_ID']}/{item}")
        if any(not str(row.get(field) or '').strip() for field in
               ('Display_Name', 'Related_Tissues', 'Primary_Item_IDs', 'Source_Reference')) or row.get('Review_Status') not in {'待临床审核', '已临床审核'}:
            errors.append(f"Missing location content: {row['Area_ID']}")
    guide_ids = {row['Guide_ID'] for row in extra['ConsultationGuide']}
    for row in extra["ConsultationGuide"]:
        required_fields = ("Guide_ID", "Topic", "Core_Prompt", "Question_Purpose", "Information_Gained", "Assessment_Help", "Source_Reference")
        if any(not str(row.get(field) or "").strip() for field in required_fields):
            errors.append(f"Missing consultation content: {row.get('Guide_ID')}")
        if row.get("Input_Mode") != "Display_Only":
            errors.append(f"Consultation cannot collect answers: {row['Guide_ID']}")
        if row.get("Scope") not in {"GENERAL", "REGION", "LOCATION"}:
            errors.append(f"Invalid consultation scope: {row['Guide_ID']}")
        scope_regions = set(str(row.get("Region_IDs") or "").split("|")) - {""}
        scope_modules = set(str(row.get("Module_Codes") or "").split("|")) - {""}
        if not scope_regions <= region_ids or not scope_modules or not scope_modules <= module_ids:
            errors.append(f"Unknown consultation region/module: {row['Guide_ID']}")
        if row.get("Scope") == "REGION" and not scope_regions:
            errors.append(f"Regional consultation lacks regions: {row['Guide_ID']}")
        scope_areas = set(str(row.get('Area_IDs') or '').split('|')) - {''}
        if not scope_areas <= area_ids or (row.get('Scope') == 'LOCATION' and not scope_areas):
            errors.append(f"Invalid consultation location: {row['Guide_ID']}")
        if row.get('Replaces_Guide_ID') and row['Replaces_Guide_ID'] not in guide_ids:
            errors.append(f"Unknown replaced consultation: {row['Guide_ID']}")
        if not isinstance(row.get("Display_Order"), (int, float)) or row.get("Review_Status") not in {"待临床审核", "已临床审核"}:
            errors.append(f"Invalid consultation order/review status: {row['Guide_ID']}")
    for row in tables["07_ContextRule"]:
        if row.get("public_group") not in {"ACTIVITY", "TIMING", "SYMPTOM", "ASSOCIATED", "RISK", "INTERNAL"}:
            errors.append(f"Invalid public context group: {row['rule_id']}")
        if row.get("public_group") != "INTERNAL" and not str(row.get("public_label") or "").strip():
            errors.append(f"Missing public context label: {row['rule_id']}")
    errors += [f"Unknown clinical purpose item: {row['Item_ID']}" for row in extra["ClinicalPurpose"] if row["Item_ID"] not in item_ids]
    errors += [f"Unknown assessment map: {row['Assessment_Map_ID']}" for row in extra["FindingFeature"] if row["Assessment_Map_ID"] not in assessment_ids]
    errors += [f"Unknown feature: {row['Feature_ID']}" for row in extra["InterventionReference"] if row["Feature_ID"] not in feature_ids]
    for sheet, fields in {
        "ClinicalPurpose": ("Item_ID", "Clinical_Purpose", "Assessment_Focus"),
        # Finding_Feature is the public explanation. A second interpretation is
        # optional: requiring one encouraged redundant filler beneath clear text.
        "FindingFeature": ("Feature_ID", "Assessment_Map_ID", "Finding_Name", "Finding_Feature"),
        "InterventionReference": ("Reference_ID", "Feature_ID", "Treatment_Goal", "Treatment_Method"),
    }.items():
        for row in extra[sheet]:
            if any(not str(row.get(field) or "").strip() for field in fields):
                errors.append(f"Missing required clinical content in {sheet}: {next(iter(row.values()))}")
    for row in extra["InterventionReference"]:
        if not str(row.get("Treatment_Method") or "").strip("； "):
            errors.append(f"Empty treatment method: {row['Reference_ID']}")
    for row in extra["FindingFeature"]:
        prompt = str(row.get("Related_Location_Prompt") or "").strip()
        label = str(row.get("Related_Location_Action_Label") or "").strip()
        if bool(prompt) != bool(label):
            errors.append(f"Incomplete location review hint: {row['Feature_ID']}")
        target_regions = set(str(row.get('Related_Location_Region_IDs') or '').split('|')) - {''}
        if not target_regions <= region_ids:
            errors.append(f"Unknown location review region: {row['Feature_ID']}")
        if row.get('Related_Location_Hide_Option_When_Current') is True and (not target_regions or not prompt):
            errors.append(f"Incomplete same-region visibility rule: {row['Feature_ID']}")
    feature_map = {row["Feature_ID"]: row["Assessment_Map_ID"] for row in extra["FindingFeature"]}
    for row in extra["InterventionReference"]:
        map_id = feature_map.get(row["Feature_ID"])
        if map_id and map_id not in assessment_ids:
            errors.append(f"Treatment references invalid assessment map: {row['Reference_ID']}")
    for rows in (extra["ClinicalPurpose"], extra["FindingFeature"], extra["InterventionReference"]):
        errors += [f"Missing content/source/status: {next(iter(row.values()))}" for row in rows
                   if not row.get("Source_URL") or row.get("Review_Status") not in {"待临床审核", "已临床审核"}]
    region_modules = {row["region_id"]: row["module"] for row in tables["01_Region"]}
    knee_items = {row["clinical_item_id"] for row in tables["05_RegionClinicalMap"]
                  if region_modules.get(row["region_id"]) == "KNEE"}
    purpose_items = {row["Item_ID"] for row in extra["ClinicalPurpose"]}
    feature_maps = {row["Assessment_Map_ID"] for row in extra["FindingFeature"]}
    reference_features = {row["Feature_ID"] for row in extra["InterventionReference"]}
    errors += [f"Knee item lacks clinical purpose: {item_id}" for item_id in sorted(knee_items - purpose_items)]
    errors += [f"Knee item lacks a specific finding: {item_id}" for item_id in sorted(knee_items)
               if not any(row["clinical_item_id"] == item_id and row["map_id"] in feature_maps
                          for row in tables["08_AssessmentFindingMap"])]
    # A localisation or unexplained swelling finding can legitimately have no
    # treatment reference. Requiring one encouraged unsupported generic methods.
    errors += [f"Broad knee treatment placeholder: {row['Reference_ID']}"
               for row in extra["InterventionReference"]
               if feature_map.get(row["Feature_ID"]) in feature_maps
               and assessment_by_id[feature_map[row["Feature_ID"]]]["clinical_item_id"] in knee_items
               and any(term in str(row["Treatment_Method"]) for term in
                       ("膝及相关关节活动改善", "局部症状管理"))]
    local_maps = {row["map_id"] for row in tables["08_AssessmentFindingMap"]
                  if row["clinical_item_id"].startswith("KNEE-LT-")}
    for row in extra["InterventionReference"]:
        if feature_map.get(row["Feature_ID"]) not in local_maps:
            continue
        if not row.get("Source_Note"):
            errors.append(f"Local treatment lacks evidence scope: {row['Reference_ID']}")
        if any(term in str(row["Treatment_Method"]) for term in
               ("调整诱发动作负荷", "活动与负荷调整", "肿胀管理", "局部贴扎辅助",
                "诱发任务负荷调整", "低刺激主动活动", "冷疗", "冰敷")):
            errors.append(f"Generic local tissue method: {row['Reference_ID']}")
    errors += [f"Invalid dimension guide: {row['Dimension_Key']}" for row in extra["DimensionGuide"]
               if row["Dimension_Key"] not in dimensions or not row["Public_Description"]]
    if errors:
        raise ValueError("Clinical extension validation failed:\n" + "\n".join(errors))
    labels = json.loads((ROOT / "scripts" / "region_labels.json").read_text(encoding="utf-8"))
    if {row["region_id"] for row in tables["01_Region"] if row["module"] == "KNEE"} != set(labels):
        raise ValueError("Knee region labels do not match workbook IDs")
    data = {
        "schema_version": "rehabmind-runtime-3",
        "database_version": next(row["value"] for row in tables["10_Governance"] if row["key"] == "database_version"),
        "workbook_sha256": hashlib.sha256(WORKBOOK.read_bytes()).hexdigest(),
        "clinical_extension_sha256": hashlib.sha256(CLINICAL_EXTENSION.read_bytes()).hexdigest(),
        "region_display_names": labels,
        "tables": {name: table for name, table in tables.items() if name != "09_Evidence"},
        "clinical_extension": extra,
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Generated {OUTPUT}")
    print(" ".join(f"{name}={len(rows)}" for name, rows in tables.items()))


if __name__ == "__main__":
    main()
