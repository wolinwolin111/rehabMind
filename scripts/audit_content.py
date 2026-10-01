"""Read-only inventory of text currently eligible for RehabMind display."""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "knowledge" / "RehabMind_LowerLimb_KnowledgeDB_v3.4.xlsx"
VISIBLE_FIELDS = {
    "01_Region": ("l3_region",),
    "02_ClinicalItem": ("display_name", "check_focus"),
    "04_Intervention": ("display_name",),
    "05_RegionClinicalMap": ("direction_type",),
    "07_ContextRule": ("source_context", "boundary"),
    "08_AssessmentFindingMap": ("trigger_condition",),
}
DEVELOPER_TERMS = re.compile(
    r"ClinicalItem|Finding|Intervention|ContextRule|L3|KNEE-[A-Z0-9-]+|AF-[A-Z0-9-]+|CORE-[A-Z0-9-]+|"
    r"全局规则|映射|命中|不自动|自动进入|底层|审计|治理|复用|本项目|本表|规则引擎",
    re.I,
)
SCHEMA_TERMS = re.compile(r"ClinicalItem|Finding|Safety|L3|(?:CORE|KNEE|AF|LL|THIGH)-[A-Z0-9-]+|UI分组|Resolver|Modifier", re.I)
RULE_TERMS = re.compile(r"不自动|按已确认|优先级|只作补充检查入口|只组织检查|只用于组织|只作为背景|不单独生成干预")
PROCESS_TERMS = re.compile(r"按已确认|若另有|只进入|先进入|进入.*?(路径|处理|训练)|优先级|不预设|不以.*?作为入口|不直接等同")
FAMILIAR = re.compile(r"熟悉症状")
CODE_LABEL = re.compile(r"^[A-Z][A-Z0-9_]+$")


def rows(sheet):
    iterator = sheet.iter_rows(values_only=True)
    headers = next(iterator)
    for excel_row, values in enumerate(iterator, 2):
        if values[0] is None or not str(values[0]).strip():
            continue
        yield excel_row, dict(zip(headers, values))


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    parser.add_argument("category", choices=("summary", "exposure", "developer", "process", "familiar", "long", "duplicates", "interventions", "context", "regions"))
    args = parser.parse_args()
    book = load_workbook(WORKBOOK, read_only=True, data_only=True)
    records = []
    sheet_counts = {}
    for sheet_name, fields in VISIBLE_FIELDS.items():
        sheet = book[sheet_name]
        source_rows = list(rows(sheet))
        sheet_counts[sheet_name] = len(source_rows)
        for excel_row, row in source_rows:
            for field in fields:
                value = str(row.get(field) or "").strip()
                if value:
                    records.append((sheet_name, excel_row, str(row[next(iter(row))]), field, value))
    hits = {
        "developer": [r for r in records if DEVELOPER_TERMS.search(r[4])],
        "process": [r for r in records if PROCESS_TERMS.search(r[4])],
        "familiar": [r for r in records if FAMILIAR.search(r[4])],
        "long": [r for r in records if len(r[4]) > 100],
        "regions": [r for r in records if r[0] == "01_Region" and CODE_LABEL.fullmatch(r[4])],
    }
    if args.category == "exposure":
        labels = json.loads((ROOT / "scripts" / "region_labels.json").read_text(encoding="utf-8"))
        all_items = {row["item_id"] for _, row in rows(book["02_ClinicalItem"])}
        item_ids = {row["clinical_item_id"] for _, row in rows(book["05_RegionClinicalMap"])}
        context_rows = list(rows(book["07_ContextRule"]))
        module_prefixes = ("KNEE-", "AF-", "LL-", "THIGH-")
        eligible_context = [(number, row) for number, row in context_rows if row["rule_type"] != "PROFILE" and str(row["rule_id"]).startswith(module_prefixes)]
        for _, row in eligible_context:
            target = str(row.get("target") or "")
            item_ids.update(item for item in all_items if re.search(r"(?<![A-Za-z0-9-])" + re.escape(item) + r"(?![A-Za-z0-9-])", target))
        mapped_interventions = {row["intervention_id"] for _, row in rows(book["06_FindingInterventionMap"])}
        exposed = []
        for record in records:
            sheet, _, identifier, field, _ = record
            if sheet == "01_Region" and identifier in labels:
                continue
            if sheet == "02_ClinicalItem" and identifier not in item_ids:
                continue
            if sheet == "04_Intervention" and identifier not in mapped_interventions:
                continue
            if sheet == "08_AssessmentFindingMap" and identifier not in {row["map_id"] for _, row in rows(book["08_AssessmentFindingMap"]) if row["clinical_item_id"] in item_ids}:
                continue
            if sheet == "07_ContextRule":
                context = next((row for _, row in eligible_context if row["rule_id"] == identifier), None)
                if context is None:
                    continue
                target = str(context.get("target") or "")
                has_item = any(re.search(r"(?<![A-Za-z0-9-])" + re.escape(item) + r"(?![A-Za-z0-9-])", target) for item in all_items)
                if field == "boundary" and context["rule_type"] != "SAFETY" and has_item:
                    continue
            exposed.append(record)
        for number, row in context_rows:
            if row["rule_type"] == "PROFILE":
                exposed.append(("07_ContextRule", number, str(row["rule_id"]), "target", str(row["target"])))
        print("potentially_rendered_cells", len(exposed), "distinct_records", len({(r[0], r[2]) for r in exposed}))
        print("by_sheet_field", dict(sorted(Counter((r[0], r[3]) for r in exposed).items())))
        language = [r for r in exposed if SCHEMA_TERMS.search(r[4]) or RULE_TERMS.search(r[4])]
        familiar_variants = [r for r in exposed if re.search(r"熟悉.{0,3}症状", r[4])]
        print("schema_or_rule_union", len(language), "by_sheet", dict(sorted(Counter(r[0] for r in language).items())))
        print("familiar_variants", len(familiar_variants), "by_sheet", dict(sorted(Counter(r[0] for r in familiar_variants).items())))
        for name, pattern in (("schema", SCHEMA_TERMS), ("rule_language", RULE_TERMS), ("developer_candidates", DEVELOPER_TERMS), ("process", PROCESS_TERMS), ("familiar_exact", FAMILIAR)):
            matched = [r for r in exposed if pattern.search(r[4])]
            print(name, "cells", len(matched), "records", len({(r[0], r[2]) for r in matched}), "unique_text", len({r[4] for r in matched}))
            if name in {"schema", "rule_language"}:
                for record in matched:
                    print(*record, sep=" | ")
        return
    if args.category == "summary":
        print("source_rows", sheet_counts)
        print("nonempty_visible_cells", len(records))
        for name, group in hits.items():
            print(name, "cells", len(group), "unique_records", len({(r[0], r[2]) for r in group}), "unique_text", len({r[4] for r in group}))
        for sheet_name, fields in VISIBLE_FIELDS.items():
            values = [r for r in records if r[0] == sheet_name]
            print(sheet_name, "visible_cells", len(values), "developer", sum(r in hits["developer"] for r in values), "process", sum(r in hits["process"] for r in values), "familiar", sum(r in hits["familiar"] for r in values), "long", sum(r in hits["long"] for r in values))
        return
    if args.category == "duplicates":
        for sheet_name, field in (("08_AssessmentFindingMap", "trigger_condition"), ("02_ClinicalItem", "check_focus"), ("04_Intervention", "display_name")):
            groups = defaultdict(list)
            for record in records:
                if record[0] == sheet_name and record[3] == field:
                    groups[record[4]].append(record[2])
            for value, ids in sorted(groups.items(), key=lambda pair: -len(pair[1])):
                if len(ids) > 1:
                    print(sheet_name, field, len(ids), ",".join(ids), value)
        return
    if args.category == "interventions":
        for record in records:
            if record[0] == "04_Intervention":
                print(*record, sep=" | ")
        return
    if args.category == "context":
        for record in records:
            if record[0] == "07_ContextRule":
                print(*record, sep=" | ")
        return
    for record in hits[args.category]:
        print(*record, sep=" | ")


if __name__ == "__main__":
    main()
