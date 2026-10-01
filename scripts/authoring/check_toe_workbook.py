"""Read-only preservation check against the pre-edit workbook."""
import json
from pathlib import Path
from openpyxl import load_workbook
root=Path(__file__).resolve().parents[2]
before=load_workbook(root/'build/authoring/backups/toe-locations-20261001/RehabMind_ClinicalExtension_v1.1.xlsx',data_only=False)
after=load_workbook(root/'knowledge/RehabMind_ClinicalExtension_v1.1.xlsx',data_only=False)
assert before.sheetnames==after.sheetnames
for name in before.sheetnames:
    a,b=before[name],after[name]
    for row in a:
        for cell in row:
            other=b.cell(cell.row,cell.column)
            assert cell.value==other.value,(name,cell.coordinate,cell.value,other.value)
            assert cell.number_format==other.number_format,(name,cell.coordinate,'format')
    if name!='LocationGuide':
        assert a.max_row==b.max_row,(name,'rows')
    else:
        assert b.max_row==a.max_row+2
spec=json.loads((root/'knowledge/review/toe_locations_v1.json').read_text(encoding='utf-8'))
rows=list(after['LocationGuide'].values)
records={r[0]:dict(zip(rows[0],r)) for r in rows[1:] if r[0]}
for key,value in spec['extension']['LocationGuide'].items():
    assert records[key]==value
print('All prior cells and number formats preserved; two complete toe guide records added.')
