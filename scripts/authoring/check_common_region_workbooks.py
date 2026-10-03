"""Read-only cell/style comparison against the current batch's source copies."""
import json
from copy import copy
from pathlib import Path
import sys
from openpyxl import load_workbook

root=Path(__file__).resolve().parents[2]
batch=4 if sys.argv[1]=='--all' else int(sys.argv[1])
specs=json.loads((root/'knowledge/review/common_region_v1.json').read_text(encoding='utf-8'))['batches']
spec=next(b for b in specs if b['batch']==batch)
if sys.argv[1]=='--all':
    spec={'main':{},'extension':{}}
    for b in sorted(specs,key=lambda b:b['batch']):
        for section in ['main','extension']:
            for name,rows in b[section].items():
                for key,patch in rows.items():spec[section].setdefault(name,{}).setdefault(key,{}).update(patch)
report=[]
for section,filename in [('main','RehabMind_KnowledgeDB_v3.8.xlsx'),('extension','RehabMind_ClinicalExtension_v1.1.xlsx')]:
    before=load_workbook(root/f'build/authoring/backups/common-regions-batch{batch}-20261002'/filename)
    after=load_workbook(root/'knowledge'/filename)
    assert before.sheetnames==after.sheetnames
    changed,new=0,0
    for name in before.sheetnames:
        a,b=before[name],after[name]
        assert a.freeze_panes==b.freeze_panes and str(a.merged_cells)==str(b.merged_cells),name
        assert list(a.tables)==list(b.tables),name
        assert str(a.data_validations)==str(b.data_validations),name
        headers=[c.value for c in a[1]]
        permitted=spec[section].get(name,{})
        for row in a:
            key=row[0].value
            for cell in row:
                actual=b.cell(cell.row,cell.column)
                if cell.value!=actual.value:
                    field=headers[cell.column-1]
                    assert field in permitted.get(key,{}) and permitted[key][field]==actual.value,(name,cell.coordinate)
                    changed+=1
                assert copy(cell.font)==copy(actual.font) and copy(cell.fill)==copy(actual.fill) and copy(cell.alignment)==copy(actual.alignment) and cell.number_format==actual.number_format,(name,cell.coordinate,'style')
        for row in b.iter_rows(min_row=a.max_row+1):
            if not row[0].value:continue
            key=row[0].value
            assert key in permitted,(name,key)
            for field,value in zip(headers,(c.value for c in row)):
                assert value==permitted[key].get(field),(name,key,field)
            new+=1
        for key,patch in permitted.items():
            row=next(r for r in b if r[0].value==key)
            for field,value in patch.items():assert row[headers.index(field)].value==value,(key,field)
    report.append({'workbook':filename,'changed_existing_cells':changed,'new_rows':new,'unrelated_values_and_styles_preserved':True})
suffix='all' if sys.argv[1]=='--all' else f'batch{batch}'
(root/f'knowledge/review/common_region_{suffix}_workbook_verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(report,ensure_ascii=False))
