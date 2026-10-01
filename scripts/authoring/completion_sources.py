"""Read authorized records and check workbook preservation; never author Excel."""
import json, pathlib, re, hashlib, sys
root = pathlib.Path(__file__).resolve().parents[2]
index = json.loads((root/'knowledge/review/local_record_sources.private.json').read_text(encoding='utf-8'))
pattern = re.compile(r'股直肌|股外侧肌|股内侧肌|股中间肌|股二头肌|半腱肌|半膜肌|腘绳肌|缝匠肌|股薄肌|内收肌|阔筋膜张肌|髂腰肌|髌骨.*松动|腓骨头.*松动|跖趾.*松动|第一.*松动|胫骨.*旋转|TENS|电刺激')
if '--display-verify' in sys.argv:
    from openpyxl import load_workbook
    spec = json.loads((root/'knowledge/review/lower_limb_display_clarity_v1.json').read_text(encoding='utf-8'))
    changed_regions = {r['region_id'] for r in spec['regions']}
    added = {'01_Region':['related_tissues','related_tissues_source'], 'FindingFeature':['Related_Location_Region_IDs','Related_Location_Hide_Option_When_Current']}
    for filename in ['RehabMind_KnowledgeDB_v3.8.xlsx','RehabMind_ClinicalExtension_v1.1.xlsx']:
        before = load_workbook(root/'build/authoring/backups/display-clarity-20261001'/filename)
        after = load_workbook(root/'knowledge'/filename)
        assert before.sheetnames == after.sheetnames
        for name in before.sheetnames:
            left,right = before[name],after[name]
            assert left.freeze_panes == right.freeze_panes
            assert list(right.values)[0] == tuple(list(list(left.values)[0]) + added.get(name,[])),name
            assert left.max_row == right.max_row,name
            for row in left:
                for cell in row:
                    other = right.cell(cell.row,cell.column)
                    assert cell.value == other.value and cell._style == other._style,(name,cell.coordinate)
                if name == '01_Region' and row[0].value in changed_regions: continue
                assert left.row_dimensions[row[0].row].height == right.row_dimensions[row[0].row].height
            for col,dim in left.column_dimensions.items(): assert dim.width == right.column_dimensions[col].width,(name,col)
        before.close();after.close()
    print(json.dumps({'existing_cells_and_styles_preserved':True,'region_tissues_added':len(spec['regions']),'visibility_rules_added':len(spec['navigation'])}))
elif '--snapshot' in sys.argv:
    from openpyxl import load_workbook
    runtime = json.loads((root/'build/knowledge/runtime.json').read_text(encoding='utf-8'))
    for filename, target in [('RehabMind_KnowledgeDB_v3.8.xlsx',runtime['tables']),('RehabMind_ClinicalExtension_v1.1.xlsx',runtime['clinical_extension'])]:
        book = load_workbook(root/'build/authoring/backups/lower-limb-completion-20261001'/filename,read_only=True,data_only=True)
        for name in book.sheetnames:
            if name not in target: continue
            rows = iter(book[name].values)
            headers = next(rows)
            target[name] = [dict(zip(headers,row)) for row in rows if row[0] is not None]
        book.close()
    snapshot = root/'build/authoring/completion-original-runtime.json'
    if not snapshot.exists(): snapshot.write_text(json.dumps(runtime,ensure_ascii=False),encoding='utf-8')
    print('Original workbook content snapshot preserved.')
elif '--verify' not in sys.argv:
    hits = []
    for source in index['files']:
        lines = (pathlib.Path(index['root'])/source['file']).read_text(encoding=source['encoding']).splitlines()
        for n, line in enumerate(lines,1):
            match = pattern.search(line)
            if match:
                hits.append({'source_id':source['id'],'line':n,'fragment':line[max(0,match.start()-5):match.end()+85].strip()})
    (root/'build/authoring/completion-record-evidence.json').write_text(json.dumps(hits,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(hits[:90],ensure_ascii=False))
    print('MATCH_COUNT',len(hits))
else:
    from openpyxl import load_workbook
    spec = json.loads((root/'knowledge/review/lower_limb_completion_v1.json').read_text(encoding='utf-8'))
    fields={'LocationGuide':'locations','ConsultationGuide':'guides','ClinicalPurpose':'purposes','FindingFeature':'features','InterventionReference':'references'}
    allowed={sheet:{r[next(iter(r))] for r in spec[field]} for sheet,field in fields.items()}
    allowed.update({sheet:{r[next(iter(r))] for r in rows} for sheet,rows in spec['main_updates'].items()})
    display_spec = json.loads((root/'knowledge/review/lower_limb_display_clarity_v1.json').read_text(encoding='utf-8'))
    allowed.setdefault('01_Region',set()).update(r['region_id'] for r in display_spec['regions'])
    added_headers = {'01_Region':['related_tissues','related_tissues_source'], 'FindingFeature':['Related_Location_Region_IDs','Related_Location_Hide_Option_When_Current']}
    changes=[]
    for filename in ['RehabMind_ClinicalExtension_v1.1.xlsx','RehabMind_KnowledgeDB_v3.8.xlsx']:
        before=load_workbook(root/'build/authoring/backups/lower-limb-completion-20261001'/filename)
        after=load_workbook(root/'knowledge'/filename)
        assert before.sheetnames==after.sheetnames
        for name in before.sheetnames:
            left,right=before[name],after[name]
            assert left.freeze_panes==right.freeze_panes
            assert tuple(list(list(left.values)[0])+added_headers.get(name,[]))==list(right.values)[0],name
            source={r[0]:r for r in left.values if r[0]}
            target={r[0]:tuple(r[:left.max_column]) for r in right.values if r[0]}
            retired=set(spec.get('retired_features',[])) if name=='FindingFeature' else set(spec.get('retired_references',[])) if name=='InterventionReference' else set()
            assert set(source)-set(target)==retired,(name,'removed')
            assert set(target)-set(source)<=allowed.get(name,set()),(name,'added')
            target_rows={r[0].value:r for r in right if r[0].value}
            for r in left:
                key=r[0].value
                if key is None or key in retired: continue
                if source[key]!=target[key]:
                    assert key in allowed.get(name,set()),(name,key,'value')
                    changes.append({'sheet':name,'id':key})
                if key in allowed.get(name,set()):continue
                dest=target_rows[key]
                for cell,other in zip(r,dest):
                    assert cell.value==other.value and cell._style==other._style,(name,key,'preservation')
                assert left.row_dimensions[r[0].row].height==right.row_dimensions[dest[0].row].height
            for col,dim in left.column_dimensions.items(): assert dim.width==right.column_dimensions[col].width,(name,col)
        before.close();after.close()
    cited=set(re.findall(r'local-records://(LR-[0-9a-f]+)',json.dumps(spec)))
    by_id={r['id']:r for r in index['files']}
    for sid in cited:
        source=by_id[sid]
        assert hashlib.sha256((pathlib.Path(index['root'])/source['file']).read_bytes()).hexdigest()==source['sha256']
    report={'unrelated_values_styles_preserved':True,'changed_records':changes,'verified_local_sources':len(cited)}
    (root/'build/authoring/completion-preservation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'unrelated_values_styles_preserved':True,'changed_records':len(changes),'verified_local_sources':len(cited)}))
