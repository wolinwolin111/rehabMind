"""Read-only comparison of saved source workbooks and their pre-edit copies."""
import json
from copy import copy
from pathlib import Path
from openpyxl import load_workbook

root=Path(__file__).resolve().parents[2]
spec=json.loads((root/'knowledge/review/region_calibration_v1.json').read_text(encoding='utf-8'))
allowed={'main':{},'extension':{}}
for batch in spec['batches']:
    for section in allowed:
        for sheet,records in batch[section].items():
            for key,patch in records.items():
                allowed[section].setdefault(sheet,{}).setdefault(key,{}).update(patch)
report=[]
for section,filename in [('main','RehabMind_KnowledgeDB_v3.8.xlsx'),('extension','RehabMind_ClinicalExtension_v1.1.xlsx')]:
    before=load_workbook(root/'build/authoring/backups/region-calibration-20261002'/filename)
    after=load_workbook(root/'knowledge'/filename)
    assert before.sheetnames==after.sheetnames, 'Sheet structure changed'
    modified=0
    for name in before.sheetnames:
        a,b=before[name],after[name]
        assert a.freeze_panes==b.freeze_panes, f'Freeze panes changed: {name}'
        assert str(a.merged_cells)==str(b.merged_cells), f'Merges changed: {name}'
        assert len(a.tables)==len(b.tables), f'Tables changed: {name}'
        assert len(a.data_validations.dataValidation)==len(b.data_validations.dataValidation), f'Validation count changed: {name}'
        headers=[cell.value for cell in a[1]]
        for row in a:
            record_id=a.cell(row[0].row,1).value
            for cell in row:
                actual=b.cell(cell.row,cell.column)
                if cell.value!=actual.value:
                    field=headers[cell.column-1]
                    patch=allowed[section].get(name,{}).get(record_id,{})
                    assert field in patch and patch[field]==actual.value, f'Unexpected value change: {name}/{cell.coordinate}'
                    modified+=1
                assert copy(cell.font)==copy(actual.font) and copy(cell.fill)==copy(actual.fill) and copy(cell.alignment)==copy(actual.alignment) and cell.number_format==actual.number_format, f'Existing style changed: {name}/{cell.coordinate}'
        new_rows=[row for row in b.iter_rows(min_row=a.max_row+1) if row[0].value] if b.max_row>a.max_row else []
        for row in new_rows:
            assert row[0].value in allowed[section].get(name,{}), f'Unexpected new row: {name}/{row[0].value}'
    report.append({'workbook':filename,'modified_existing_cells':modified,'existing_styles_preserved':True,'sheet_structure_preserved':True})
(root/'knowledge/review/region_calibration_workbook_verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(report,ensure_ascii=False))
